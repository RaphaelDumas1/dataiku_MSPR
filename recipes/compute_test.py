import dataiku
import pandas as pd
import openpyxl
import time
from io import BytesIO

folder_id = "aPmnwurD"
files_names = ["MSPR - Demographie.xlsx"]
exclude_sheets = ["Quotient familiale"]


def run(folder_id, files_names):
    client = dataiku.api_client()
    project = client.get_project("MSPR")

    folder = dataiku.Folder(folder_id, project_key=project.project_key)
    folder_paths = [name.strip() for name in files_names.split(";") if name.strip()]

    # List the datasets in the project
    datasets_in_project = []
    for dataset in project.list_datasets():
        datasets_in_project.append(dataset.get('name'))

    for file_index, file_name in enumerate(folder_paths):
        with folder.get_download_stream(file_name) as file_handle:
            try:
                ss = openpyxl.load_workbook(BytesIO(file_handle.read()))
            except Exception:
                continue

        for sheet in ss.sheetnames:
             if sheet in exclude_sheets:
                continue

            ss_sheet = ss[sheet]
            title = ss_sheet.title

            title = '_'.join(title.split()).replace(')', '').replace('(', '').replace('/', '_').replace('.', '_')

            for sheet in ss.sheetnames:
                ss_sheet = ss[sheet]
                title = ss_sheet.title

                title = '_'.join(title.split()).replace(')', '').replace('(', '').replace('/', '_').replace('.', '_')

                data = list(ss_sheet.values)
                headers = data[0]
                rows = data[1:]
                df = pd.DataFrame(rows, columns=headers)
                test = dataiku.Dataset(title)
                test.write_with_schema(df)

for name in files_names:
    run(folder_id, name)
